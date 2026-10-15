"""Synthetic local PPTX production and renderer-backed evidence for E2E7."""

from __future__ import annotations

import hashlib
import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree


class PresentationMockError(ValueError):
    pass


def build_deck(brief: dict) -> bytes:
    """Create three editable slides containing only approved synthetic labels."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    if brief.get("worker") != "mock" or brief.get("slides") != 3:
        raise PresentationMockError("unsupported_mock_brief")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    headings = ("F.R.I.D.A.Y.", "Arbeitsfluss", "Prüfgrenzen")
    bodies = (
        "Synthetischer Überblick für interne Tests",
        "Brief → Entwurf → Datei → Render → Prüfung",
        "Mockdaten · lokale Datei · keine Office-Animation",
    )
    for number, (heading, body) in enumerate(zip(headings, bodies, strict=True), start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(15, 25, 41)
        tile = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.7), Inches(4.8))
        tile.fill.solid()
        tile.fill.fore_color.rgb = RGBColor(28, 46, 68)
        tile.line.fill.background()
        for y, text, size, color in ((2.0, heading, 40, RGBColor(255, 255, 255)),
                                     (3.25, body, 22, RGBColor(143, 223, 208))):
            shape = slide.shapes.add_textbox(Inches(1.45), Inches(y), Inches(10.3), Inches(1.3))
            paragraph = shape.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = color
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.1), Inches(5.2), Inches(0.5), Inches(0.5))
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(143, 223, 208)
        marker.line.fill.background()
        marker.text = str(number)
    output = io.BytesIO()
    deck.save(output)
    validate_deck(output.getvalue(), expected_slides=3)
    return output.getvalue()


def validate_deck(data: bytes, *, expected_slides: int) -> None:
    """Check the package, slide XML and actual parseability, not a fixture flag."""
    from pptx import Presentation

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if archive.testzip() is not None:
                raise PresentationMockError("corrupt_pptx_zip")
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "ppt/presentation.xml"}
            required |= {f"ppt/slides/slide{index}.xml" for index in range(1, expected_slides + 1)}
            if not required <= names:
                raise PresentationMockError("missing_pptx_parts")
            for name in required:
                ElementTree.fromstring(archive.read(name))
        deck = Presentation(io.BytesIO(data))
        if len(deck.slides) != expected_slides or any(not any(shape.has_text_frame and shape.text.strip() for shape in slide.shapes) for slide in deck.slides):
            raise PresentationMockError("invalid_slide_count_or_content")
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise PresentationMockError("invalid_pptx") from exc


def find_renderer() -> Path | None:
    for name in ("libreoffice", "soffice"):
        executable = shutil.which(name)
        if executable:
            return Path(executable)
    for candidate in (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ):
        if candidate.is_file():
            return candidate
    return None


def render_deck(pptx_path: Path, *, renderer: Path | None = None) -> dict:
    """Render all slides through local Office CLI, then rasterize PDF pages."""
    renderer = renderer or find_renderer()
    if renderer is None:
        raise PresentationMockError("renderer_unavailable")
    try:
        import fitz
    except ModuleNotFoundError as exc:
        raise PresentationMockError("pdf_rasterizer_unavailable") from exc

    pptx_path = pptx_path.resolve()
    profile = pptx_path.parent / "office-profile"
    profile.mkdir(exist_ok=True)
    pdf_path = pptx_path.with_suffix(".pdf")
    command = [str(renderer), f"-env:UserInstallation={profile.as_uri()}", "--headless",
               "--convert-to", "pdf", "--outdir", str(pptx_path.parent), str(pptx_path)]
    try:
        completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise PresentationMockError("render_failed") from exc
    if completed.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise PresentationMockError("render_failed")
    images = []
    with fitz.open(pdf_path) as pdf:
        if len(pdf) != 3:
            raise PresentationMockError("render_page_count_mismatch")
        for index, page in enumerate(pdf, start=1):
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            samples = pix.samples
            stride = max(3, (len(samples) // 12000 // 3) * 3)
            colors = {samples[offset:offset + 3] for offset in range(0, len(samples) - 2, stride)}
            if pix.width < 500 or pix.height < 300 or not page.get_text().strip() or len(colors) <= 1:
                raise PresentationMockError("blank_or_invalid_render")
            image_path = pptx_path.with_name(f"{pptx_path.stem}-slide-{index}.png")
            pix.save(image_path)
            images.append(str(image_path))
    return {"pdf_path": str(pdf_path), "slide_images": images, "renderer": renderer.name,
            "pdf_sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest()}


def run_mock(brief: dict, *, output_root: Path | None = None, renderer: Path | None = None) -> dict:
    output_root = output_root or Path(__file__).resolve().parents[1] / "artifacts"
    output_root.mkdir(parents=True, exist_ok=True)
    directory = Path(tempfile.mkdtemp(prefix="presentation-mock-", dir=output_root))
    try:
        deck_bytes = build_deck(brief)
    except ModuleNotFoundError as exc:
        if exc.name == "pptx":
            return {"status": "partial", "reason": "pptx_builder_unavailable", "slides": 0}
        raise
    deck_path = directory / "friday-mock.pptx"
    deck_path.write_bytes(deck_bytes)
    result = {"status": "partial", "slides": 3, "deck_path": str(deck_path),
              "deck_sha256": hashlib.sha256(deck_bytes).hexdigest(), "rendered": False}
    try:
        evidence = render_deck(deck_path, renderer=renderer)
    except PresentationMockError as exc:
        result["reason"] = str(exc)
        return result
    result.update(evidence, status="completed", rendered=True)
    return result
