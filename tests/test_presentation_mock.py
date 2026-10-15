import tempfile
import unittest
import os
from pathlib import Path
from unittest.mock import patch

from adapters.presentation_mock import (
    PresentationMockError,
    build_deck,
    run_mock,
    validate_deck,
)


BRIEF = {"topic": "F.R.I.D.A.Y. Überblick", "audience": "intern", "slides": 3, "worker": "mock"}


class PresentationMockTests(unittest.TestCase):
    def test_builds_parseable_three_slide_pptx(self):
        data = build_deck(BRIEF)
        validate_deck(data, expected_slides=3)
        from pptx import Presentation
        from io import BytesIO

        slides = Presentation(BytesIO(data)).slides
        self.assertEqual(len(slides), 3)
        self.assertIn("F.R.I.D.A.Y.", " ".join(shape.text for shape in slides[0].shapes if shape.has_text_frame))
        with self.assertRaises(PresentationMockError):
            validate_deck(data[:-100], expected_slides=3)

    def test_missing_renderer_stays_partial_with_real_pptx(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("adapters.presentation_mock.find_renderer", return_value=None):
                result = run_mock(BRIEF, output_root=Path(directory))
            self.assertEqual(result["status"], "partial")
            self.assertEqual(result["reason"], "renderer_unavailable")
            self.assertFalse(result["rendered"])
            self.assertEqual(result["slides"], 3)
            validate_deck(Path(result["deck_path"]).read_bytes(), expected_slides=3)

    def test_invalid_brief_cannot_report_completed(self):
        with self.assertRaises(PresentationMockError):
            build_deck({**BRIEF, "slides": 4})

    def test_renderer_failure_cannot_report_completed(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("adapters.presentation_mock.render_deck", side_effect=PresentationMockError("render_failed")):
                result = run_mock(BRIEF, output_root=Path(directory))
            self.assertEqual(result["status"], "partial")
            self.assertEqual(result["reason"], "render_failed")
            self.assertFalse(result["rendered"])

    def test_real_renderer_when_ci_requires_it(self):
        if os.environ.get("FRIDAY_REQUIRE_PRESENTATION_RENDER") != "1":
            self.skipTest("local renderer check is opt-in; Ubuntu CI requires it")
        with tempfile.TemporaryDirectory() as directory:
            result = run_mock(BRIEF, output_root=Path(directory))
            self.assertEqual(result["status"], "completed", result)
            self.assertEqual(len(result["slide_images"]), 3)
            self.assertTrue(all(Path(path).is_file() for path in result["slide_images"]))


if __name__ == "__main__":
    unittest.main()
